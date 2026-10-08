"""Clinician review packs (100 questions each) for the checked u74-g01..g13 adaptations. Batch-1 layout, Batch-1 files as templates.
usage: python build_pack_u74.py <work dir> <template.docx> <template.xlsx> <out_dir>"""
import sys, json, copy, os, glob, re
import docx, openpyxl
from docx.oxml.ns import qn

WORK, TDOCX, TXLSX, OUT = sys.argv[1:5]
RETURN_TO = os.environ.get("RETURN_TO", "the AMaC owner")   # the 8 Oct packs were built with the owner's first name here
ld = lambda p: json.load(open(p, encoding="utf-8"))
GROUPS = sorted(glob.glob(os.path.join(WORK, "u74-g*")))
keep, retire, review = [], [], {}
for g in GROUPS:
    plan = {p["id"]: p for p in ld(g + "/ctx/plan.json")}
    revs = {r["id"]: r for r in ld(g + "/rev/C.json")}
    for q in ld(g + "/out/final.json"):
        q = dict(q); q["domain"] = plan[q["id"]]["domain"]; keep.append(q)
    for r in ld(g + "/ctx/retire_prescreen.json"):
        retire.append((r["id"], r["source_akt_id"], plan[r["id"]]["domain"], ", ".join(r["matching_ids"]),
                       "Automatic pre-screen: same scenario and answer as the live question(s) named."))
    drafts = {q["id"]: q for q in ld(g + "/draft.json")}
    for i, r in revs.items():
        review[i] = r
        if r["verdict"] == "drop":
            m = [x for x in dict.fromkeys(re.findall(r"\b[A-Z]{1,6}[0-9]{2,6}\b", r["issues"])) if x != i]
            retire.append((i, drafts[i]["source_akt_id"], plan[i]["domain"], ", ".join(m), "Checker: " + r["issues"]))
TWINS = ld(os.path.join(WORK, "twins.json")) if os.path.exists(os.path.join(WORK, "twins.json")) else {}
order = []
for q in keep:
    if q["domain"] not in order: order.append(q["domain"])
keep.sort(key=lambda q: order.index(q["domain"]))
packs = [keep[i:i + 100] for i in range(0, len(keep), 100)]
NT = len(packs)
ret_pack = {}   # each domain's retirements split across packs in proportion to that domain's kept questions per pack
for dom in {r[2] for r in retire}:
    rs = [r for r in retire if r[2] == dom]
    cnt = [sum(1 for q in p if q["domain"] == dom) for p in packs]
    tot = sum(cnt) or 1
    if not sum(cnt): cnt = [1] + [0] * (NT - 1)
    cum, start = 0, 0
    for n, c in enumerate(cnt, 1):
        cum += c
        end = round(len(rs) * cum / tot) if sum(cnt) == tot else len(rs)
        for r in rs[start:end]: ret_pack[r[0]] = n
        start = end
summary = []

for n, pk in enumerate(packs, 1):
    NN = f"{n:02d}"
    PACK, FORM = f"AMaC_PLAB1_Adapted_{NN}_Review_Pack.docx", f"AMaC_PLAB1_Adapted_{NN}_Review_Form.xlsx"
    N = len(pk)
    ret = [r for r in retire if ret_pack[r[0]] == n]
    NR = len(ret)
    doms = []
    for q in pk:
        if q["domain"] not in doms: doms.append(q["domain"])
    n_fix = sum(1 for q in pk if review[q["id"]]["verdict"] == "fix")

    d = docx.Document(TDOCX)
    body = d.element.body
    tbl_proto = copy.deepcopy(d.tables[0]._tbl)
    for el in list(body):
        if el.tag != qn("w:sectPr"): body.remove(el)
    STY = {s.name: s for s in d.styles if s.type == 1}

    def para(style, runs=()):
        p = d.add_paragraph(); p.style = STY[style]
        for r in runs:
            if r == "\n": p.add_run().add_break(); continue
            p.add_run(r[1]).bold = True if r[0] else None
        return p

    def lab(style, label, text):
        return para(style, [(True, label), (False, " "), (False, text)])

    def table(header, rows):
        t = copy.deepcopy(tbl_proto)
        trs = t.findall(qn("w:tr"))
        for r in trs[2:]: t.remove(r)
        proto = trs[1]; t.remove(proto)
        def fill(tr, vals):
            tcs = tr.findall(qn("w:tc"))
            while len(tcs) < len(vals):
                tr.append(copy.deepcopy(tcs[-1])); tcs = tr.findall(qn("w:tc"))
            for tc, v in zip(tcs, vals):
                ts = tc.findall(".//" + qn("w:t")); ts[0].text = v
                for x in ts[1:]: x.text = ""
        fill(trs[0], header)
        grid = t.find(qn("w:tblGrid")); cols = grid.findall(qn("w:gridCol"))
        while len(cols) < len(header):
            grid.append(copy.deepcopy(cols[-1])); cols = grid.findall(qn("w:gridCol"))
        for vals in rows:
            r = copy.deepcopy(proto); fill(r, vals); t.append(r)
        body.insert(len(body) - 1, t)

    para("Heading 1", [(False, f"AMaC PLAB 1 — new adaptations, review pack {NN} of {NT:02d}: clinical review pack")])
    para("First Paragraph", [(True, "For:"), (False, " the AMaC clinical reviewer "),
                             (True, "Questions:"), (False, f" {N} "),
                             (True, "Time needed:"), (False, " about 3 hours " if N > 75 else f" about {max(30, round(N * 2 / 15) * 15)} minutes "),
                             (True, "Return to:"), (False, " " + RETURN_TO + " — please do not change the website or the question bank")])
    para("Normal")
    para("Heading 2", [(False, "1. What this pack is")])
    para("First Paragraph", [(False, "These are new PLAB 1 adaptations of UKMLA AKT questions. Each one is written to replace a verbatim AKT copy that was taken off the live PLAB 1 bank. None of them is live. AMaC’s rule is "),
                             (True, "nothing goes live without a clinician’s verdict"),
                             (False, f", and they will go live only after your verdicts. This pack covers: {', '.join(doms)}. Questions are grouped by domain.")])
    para("Body Text", [(False, "Each question shows every field that would go live, the id of the AKT question it adapts, the writer’s notes, and what the independent checker found or changed.")])
    para("Heading 2", [(False, "2. Checks already done (please do not repeat these)")])
    lab("Compact", "Independent check:", "a second AI reviewer who did not write the questions checked every question against current UK guidance (NICE, GMC, UKHSA, BSG, NHS England, the NHS website and others; NICE CKS and the BNF were blocked from the checkers’ location, so they are not cited). It added a named source to every answer and fixed what it found. "
        f"{n_fix} of {N} were edited. Where the checker made a clinical correction or was unsure, this is shown under the question as “Checker”.")
    lab("Compact", "Repeats:", "each question was compared with the live PLAB 1 bank (including the one-line BX recall items), with the 2,431 held adaptations awaiting review, and with the rest of this set. Repeats are listed in section 3 instead of being shown as questions.")
    lab("Compact", "Automatic format checks (all passed):", "five options; correct option never the longest; answer line matches the key; why-wrong covers every wrong letter; a Source line; the stem differs from the AKT copy it replaces.")
    lab("Compact", "Not done:", "no human clinician has seen these. That is what this review is for.")
    para("Heading 2", [(False, "3. Sources proposed for retirement (repeats)")])
    if ret:
        para("First Paragraph", [(False, f"Adaptations of these {NR} sources would repeat, in both scenario and learning point, a question that is already live, held, or kept elsewhere in these packs. "),
                                 (True, "We propose retiring the parked source"), (False, " and adding no new question. Please tick in the form if you agree, or name any you would keep in the comments.")])
        table(["Draft id", "AKT source", "Matching ids", "Reason"], [[i, s, m, why] for i, s, dm, m, why in ret])
    else:
        para("First Paragraph", [(False, "None in this pack.")])
    para("Heading 2", [(False, "4. What to check for each question")])
    for label, text in [("Keyed answer", "— the single best answer; no other option defensible."),
                        ("Accuracy and currency", "— every statement correct today by UK guidance."),
                        ("Distractors", "plausible but clearly wrong; nothing gives the answer away."),
                        ("PLAB fit", "— right for an IMG entering UK practice; no unnecessary figures; the key does not depend on recalling a movable number.")]:
        lab("Compact", label, text)
    para("First Paragraph", [(False, "Please record verdicts in "), (True, FORM), (False, ".")])
    para("Normal")
    para("Heading 2", [(False, f"5. The {N} questions")])
    cur = None
    for q in pk:
        if q["domain"] != cur:
            cur = q["domain"]
            para("Body Text", [(True, f"{cur} — {sum(1 for x in pk if x['domain'] == cur)} question(s)")])
        para("Heading 3", [(False, f"{q['id']} · AKT source {q['source_akt_id']} · {q['difficulty']} · {q['presentation']}")])
        vign, _, ask = q["stem"].rpartition("\n\n")
        para("First Paragraph", [(False, vign.strip())])
        para("Body Text", [(False, ask.strip())])
        for L in "ABCDE":
            para("Compact", [(True, f"{L}."), (False, " "), (False, q["options"][L])])
        lab("First Paragraph", "Answer:", q["correct_answer"])
        lab("Body Text", "Why correct:", q["why_correct"])
        lab("Body Text", "Why the others are wrong:", q["why_wrong"])
        para("Body Text", [(True, "Pearl:"), (False, " " + q["pearl"]), "\n",
                           (True, "Reasoning:"), (False, " " + q["thinking"]), "\n",
                           (True, "Exam trap:"), (False, " " + q["exam_trap"]), "\n",
                           (True, "Takeaway:"), (False, " " + q["takeaway"])])
        if str(q.get("notes", "")).strip():
            lab("Body Text", "Writer’s notes:", q["notes"])
        iss = review[q["id"]].get("issues", "").strip()
        if iss:
            lab("Body Text", "Checker:", iss)
        if q["id"] in TWINS:
            lab("Body Text", "Twin in this set:", TWINS[q["id"]])
        lab("Body Text", "Verdict:", "☐ Approve ☐ Edit ☐ Reject — ________________________")
        para("Normal")
    para("Heading 2", [(False, "6. Sign-off (internal record only — never shown on the website)")])
    para("Compact", [(True, "Reviewer:"), (False, " ____________________ "), (True, "Date:"), (False, " __________")])
    para("Compact", [(True, "Approved:"), (False, " ___ "), (True, "Edited:"), (False, " ___ "), (True, "Rejected:"), (False, " ___ "),
                     (True, "Retire the sources in section 3:"), (False, " ☐ yes ☐ no")])
    para("Compact", [(True, "Comments:")])
    d.core_properties.title = f"AMaC PLAB 1 new adaptations {NN} clinical review pack"
    d.save(os.path.join(OUT, PACK))

    wb = openpyxl.load_workbook(TXLSX)
    how = wb["How to use"]
    how["A1"] = f"AMaC PLAB 1 new adaptations, pack {NN} of {NT:02d} — reviewer feedback form"
    how["A3"] = f"1. Read each question in {PACK}."
    how["A6"] = "4. The last row asks about retiring the sources in section 3." if ret else "4. (No extra rows in this pack.)"
    how["A11"] = "PAN0000"; how["B11"] = "N0000"; how["B10"] = "AKT source"
    ws = wb["Questions"]
    ws["B1"] = "AKT source"
    proto = [copy.copy(ws.cell(2, c)._style) for c in range(1, 9)]
    height = ws.row_dimensions[2].height
    ws.delete_rows(2, ws.max_row)
    rows = [[q["id"], q["source_akt_id"], q["presentation"], q["correct_answer"], None, None, None, None] for q in pk]
    if ret:
        rows.append(["Section 3", "—", f"Retire the {NR} sources listed in section 3?", None, None, None, None, None])
    for i, r in enumerate(rows, start=2):
        for c, v in enumerate(r, start=1):
            cell = ws.cell(i, c, v); cell._style = copy.copy(proto[c - 1])
        ws.row_dimensions[i].height = height
    ws.data_validations.dataValidation[0].sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"E2:E{len(rows) + 1}")
    sm = wb["Summary"]
    for r in range(2, 5): sm.cell(r, 2).value = f"=COUNTIF(Questions!$E$2:$E${N + 1},A{r})"
    sm["B5"] = f"=COUNTBLANK(Questions!$E$2:$E${N + 1})"
    sm["B6"] = N
    wb.properties.title = f"AMaC PLAB 1 new adaptations {NN} review form"
    wb.save(os.path.join(OUT, FORM))
    summary.append((PACK, N, NR, n_fix, doms))
for s in summary: print(*s)
print("kept", len(keep), "retire", len(retire), "packs", NT)
