import json, re, collections
import docx

B = '_tools/plab-adaptation/work/id-emergencies/'
fin = json.load(open(B + 'final.json', encoding='utf-8'))
skips = [q for q in json.load(open(B + 'skips_all.json', encoding='utf-8')) if q['skip']]
drops = json.load(open(B + 'checker_drops.json', encoding='utf-8'))
retire = json.load(open(B + 'ctx/retire.json', encoding='utf-8'))
rev = json.load(open(B + 'rev_all.json', encoding='utf-8'))
swaps = json.load(open(B + 'letter_swaps.json'))
src = {x['live_plab_copy']['source_akt_id']: x['live_plab_copy']
       for x in json.load(open(B + 'ctx/source.json', encoding='utf-8'))}
hold_ids = {'PAN7706', 'PAN8052', 'PAN8624'}
dupskips = [q for q in skips if q['id'] not in hold_ids]
holds = [q for q in skips if q['id'] in hold_ids]
near = {}
for r in rev:
    near.setdefault(r['id'], []).extend(r.get('near_checked', []))
flag = {
    'PAN8458': 'Rebuilt by C2 because the first draft was a near-duplicate of live PAQ14787 (PID: sexual health clinic, new partner, CMT). The angle now is PID mimicking cystitis in an out-of-hours GP setting. Drop if still judged too close to PAQ14787.',
    'PAN1434': 'PID partner management: live partner-notification items (ID031, ID045, BX994) cover chlamydia and gonorrhoea only, so overlap is moderate. Key lengthened by C1 (was the shortest option). The partner look-back wording follows BASHH PID 2019 (current partner; other recent partners may be offered screening).',
    'PAN9977': 'Closest live items are the occupational PEP questions (ID024, PAQ11795, BX490); this one is receptive anal sex with a partner of unknown status. No dose, duration or 72-hour figures used.',
    'PAN2488': 'No UK page covering hepatitis E in pregnancy could be fetched; cited to the British Liver Trust and standard teaching, with the UK genotype 3 nuance noted by C2.',
    'PAN3587': 'Standard UK teaching; no guideline page fetched.',
    'PAN3827': 'Standard UK teaching; no guideline page fetched. Distractors are other meningitis causes.',
    'PAN3783': 'Bartonella content is standard teaching; no accessible UK source (C1).',
    'PAN3224': 'Shares the chlamydia theme with PAN3756, PAN8458 and PAN1434, but tests silent infection and PID/tubal risk in a contact; C3 judged it distinct.',
}

d = docx.Document('review-packs/plab1-adaptation-ii-neurology.docx')
body = d.element.body
for el in list(body):
    if not el.tag.endswith('sectPr'):
        body.remove(el)


def P(t='', s='Normal'):
    return d.add_paragraph(t, style=s)


def lab(p, label, text):
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)


written = len(fin) + len(drops)
assert written == 27
d.add_paragraph('PLAB 1 Adaptation II: Infectious Diseases & Emergencies review pack', style='Title')
P('Review pack only. Not live, no PR, plab1-questions.js and _parked untouched. Branch plab1-adaptation-id-emergencies-review, built from main 70d7e31 (PLAB 1 bank 5,978; _parked 2,652). 17 October 2026.')
d.add_paragraph('1. Counts', style='Heading 1')
rows = [
    ('Candidates in (clean, Infectious Diseases & Emergencies)', 201),
    ('Pre-screen repeats of live questions', len(retire)),
    ('Sent to writers', 196),
    ('Writers skipped: duplicate of a live or in-batch item', len(dupskips)),
    ('Writers held: borderline duplicate, reversible', len(holds)),
    ('Written', written),
    ('Sent to checkers', written),
    ('Checker drops (duplicates)', len(drops)),
    ('IN PACK (section 3)', len(fin)),
]
assert 201 - len(retire) == 196 and 196 - len(dupskips) - len(holds) == written
t = d.add_table(rows=0, cols=2)
t.style = 'Light Grid Accent 1'
for a, b in rows:
    c = t.add_row().cells
    c[0].text = a
    c[1].text = str(b)
P()
p = P()
r = p.add_run(f'Funnel: 201 candidates -> {written} written -> {len(fin)} in pack.')
r.bold = True
L = collections.Counter(q['correct_letter'] for q in fin)
lg = sum(len(q['options'][q['correct_letter']]) == max(map(len, q['options'].values())) for q in fin)
sh = sum(len(q['options'][q['correct_letter']]) == min(map(len, q['options'].values())) for q in fin)
P('Answer letters: ' + ', '.join(f'{l} {L[l]}' for l in 'ABCDE')
  + f'. Key is the longest option in {lg} of {len(fin)}; the shortest in {sh} of {len(fin)}. '
  + f'Letters were pre-balanced in the plan (196 items); after skips and drops they were evened by swapping option positions only ({len(swaps)} items: {", ".join(swaps)}); wording unchanged.')
P('Validator (_tools/plab-adaptation/validate.py, including key longest/shortest <= 25%): OK. Live-bank checks: no ID clash and no stem >= 85% (rapidfuzz token_set_ratio, highest 68) against the 5,978 live questions; no within-pack stem pair above 68%. Content-integrity baseline unchanged at 473 warnings (no question file edited). Each checker searched the live bank for every item by learning point, scenario and keyed answer; the live ids examined are listed per item.')
P('Citations: NICE CKS pages returned 403 from this environment, so CKS-based statements are standard teaching rather than page-checked; checkers removed or reworded every citation they could not confirm. Guideline text fetched and checked: NICE NG60 (HIV testing), NG84 (sore throat), NG33 (tuberculosis), NG240 (meningitis), NG253 (sepsis), NG250 (pneumonia); UKHSA botulism, scarlet fever and diphtheria pages; NHS chlamydia, syphilis and chickenpox pages. Items with a specific concern carry a "Reviewer note".')
P('Every item is labelled "2026 UKMLA Adaptation II — adapted (Oct 2026)".')
d.add_paragraph('2. Dropped duplicates (with the live or pack ID duplicated)', style='Heading 1')
d.add_paragraph(f'2.1 Writer skips: duplicates ({len(dupskips)})', style='Heading 2')
for q in dupskips:
    P(f"{q['id']}: {q['notes'].strip()}", 'List Bullet')
d.add_paragraph(f'2.2 Checker drops ({len(drops)})', style='Heading 2')
for x in drops:
    P(f"{x['id']} ({x['checker']}): {x['issues'].strip()}", 'List Bullet')
P()
d.add_paragraph(f'3. Items proposed for go-live ({len(fin)})', style='Heading 1')
for n, q in enumerate(fin, 1):
    d.add_paragraph(f"{n}. {q['id']}: {q['presentation']}", style='Heading 2')
    P(f"Label: 2026 UKMLA Adaptation II — adapted (Oct 2026)   |   Source ID: {q['source_akt_id']}   |   Difficulty: {q['difficulty']}   |   Domain: Infectious Diseases & Emergencies")
    vig, qq = q['stem'].split('\n\n', 1)
    p = P()
    lab(p, 'Stem: ', vig)
    P(qq)
    for l in 'ABCDE':
        p = P()
        r = p.add_run(f"{l}. {q['options'][l]}")
        r.bold = (l == q['correct_letter'])
    p = P()
    lab(p, 'Key: ', q['correct_answer'])
    p = P()
    lab(p, 'Why correct: ', q['why_correct'])
    p = P()
    lab(p, 'Why others are wrong: ', q['why_wrong'])
    m = re.search(r'Source:\s*(.*)$', q['why_correct'].strip())
    p = P()
    lab(p, 'Citation: ', m.group(1) if m else '; '.join(str(s) for s in q['sources']))
    srcs = [s for s in q['sources'] if isinstance(s, str)]
    if srcs:
        p = P()
        lab(p, 'Sources opened/carried: ', '; '.join(srcs))
    nn = sorted(set(near.get(q['id'], [])))
    p = P()
    lab(p, 'Live near-matches checked: ', ', '.join(nn) if nn else 'none recorded')
    if q['id'] in flag:
        p = P()
        lab(p, 'Reviewer note: ', flag[q['id']])
    P('Verdict:  ☐ APPROVE   ☐ EDIT   ☐ REJECT      Comment: ____________________')
d.add_paragraph(f'Appendix A: pre-screen repeats ({len(retire)})', style='Heading 1')
for r in retire:
    P(f"{r['parked_id']} ({src[r['source_akt_id']]['presentation']}): live {', '.join(r['matching_ids'])}", 'List Bullet')
d.add_paragraph(f'Appendix B: holds ({len(holds)})', style='Heading 1')
P('Held as borderline duplicates: the writer judged them the same learning point as a live question that differs only in setting. A reviewer may reverse the hold. None is a confirmed duplicate.')
for q in holds:
    P(f"{q['id']}: HOLD: {q['notes'].strip()}", 'List Bullet')
d.save('review-packs/plab1-adaptation-ii-id-emergencies.docx')
