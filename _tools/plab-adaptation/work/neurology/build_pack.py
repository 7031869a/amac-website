import json,re,collections
import docx
B='_tools/plab-adaptation/work/neurology/'
fin=json.load(open(B+'final.json',encoding='utf-8'))
skips=[q for q in json.load(open(B+'skips_all.json',encoding='utf-8')) if q['skip']]
drops=json.load(open(B+'checker_drops.json',encoding='utf-8'))
retire=json.load(open(B+'ctx/retire.json',encoding='utf-8'))
src={x['live_plab_copy']['source_akt_id']:x['live_plab_copy'] for x in json.load(open(B+'ctx/source.json',encoding='utf-8'))}
pres=lambda sid:src[sid]['presentation']
hold_ids={'PAN5590','PAN8595'}
dupskips=[q for q in skips if q['id'] not in hold_ids]
holds=[q for q in skips if q['id'] in hold_ids]
flag={'PAN3658':'Close to live PAQ15793 (narcolepsy from laughter-triggered knee buckling); this one tests cataplexy against seizure, syncope and TIA. Reviewer to judge.',
 'PAN5574':'Overlaps live PAQ16161 (left MCA stroke with aphasia) and BX879; this one tests right MCA with neglect. Kept after C1 and C3 each dropped the other of the PAN5574/PAN5163 pair; PAN5163 was dropped.',
 'PAN10231':'Kept after C1 and C2 each dropped the other of the PAN10231/PAN5860/PAN10581 group (Parkinson disease non-motor features); the other two were dropped. Not edited by a checker.',
 'PAN3915':'Partly overlaps live NEU022 and NEU085; kept by C3.',
 'PAN7031':'No CKS or NHS meralgia page could be found; cited to Cochrane and standard teaching, management wording softened (C1).'}
d=docx.Document('review-packs/plab1-adaptation-ii-obstetrics-gynaecology.docx')
body=d.element.body
for el in list(body):
    if el.tag.endswith('sectPr'): continue
    body.remove(el)
P=lambda t='',s='Normal':d.add_paragraph(t,style=s)
def lab(p,label,text):
    r=p.add_run(label); r.bold=True; p.add_run(text)
written=len(fin)+len(drops); assert written==41
d.add_paragraph('PLAB 1 Adaptation II: Neurology review pack',style='Title')
P('Review pack only. Not live, no PR, plab1-questions.js untouched. Branch plab1-adaptation-neurology-review, built from main 55d577a (PLAB 1 bank 5,947). 17 October 2026.')
d.add_paragraph('1. Counts',style='Heading 1')
rows=[('Candidates in (clean, Neurology)',204),('Pre-screen repeats of live questions',len(retire)),('Sent to writers',192),
('Writers skipped: duplicate of a live or in-batch item',len(dupskips)),('Writers held: unsound learning point or bare recall',len(holds)),
('Written',written),('Sent to checkers',written),('Checker drops (duplicates)',len(drops)),('IN PACK (Appendix A)',len(fin))]
t=d.add_table(rows=0,cols=2); t.style='Light Grid Accent 1'
for a,b in rows:
    c=t.add_row().cells; c[0].text=a; c[1].text=str(b)
P()
p=P(); r=p.add_run(f'Funnel: 204 candidates -> {written} written -> {len(fin)} in pack.'); r.bold=True
L=collections.Counter(q['correct_letter'] for q in fin)
lg=sum(len(q['options'][q['correct_letter']])==max(map(len,q['options'].values())) for q in fin)
sh=sum(len(q['options'][q['correct_letter']])==min(map(len,q['options'].values())) for q in fin)
P('Answer letters: '+', '.join(f'{l} {L[l]}' for l in 'ABCDE')+f'. Key is the longest option in {lg} of {len(fin)}; the shortest in {sh} of {len(fin)}. Letters were pre-balanced in the plan (192 items), but skips and drops unbalanced the survivors, so after checking they were evened by swapping option positions only (4 items: PAN10231, PAN3658, PAN6634, PAN9746); wording unchanged.')
P('Validator (_tools/plab-adaptation/validate.py, including key longest/shortest <= 25%): OK. Live-bank checks: no ID clash and no stem >= 85% (rapidfuzz token_set_ratio, highest 72) against the 5,947 live questions; no within-pack stem pair >= 70%. Content-integrity baseline unchanged at 473 warnings (no question file edited).')
P('Citations: NICE CKS pages returned 403 from this environment, so CKS-based statements are standard teaching rather than page-checked; checkers removed or reworded every CKS citation they could not confirm. Guideline text fetched and checked by checkers: NICE NG128 (stroke), NG71 (Parkinson disease), NG127 (suspected neurological conditions), NG217 (epilepsies), NG220 (MS), NG236 (stroke rehabilitation). Items with a specific concern are flagged "Reviewer note".')
P('Every item is labelled "2026 UKMLA Adaptation II — adapted (Oct 2026)".')
d.add_paragraph('2. Dropped duplicates (with the live or pack ID duplicated)',style='Heading 1')
d.add_paragraph(f'2.1 Pre-screen repeats ({len(retire)})',style='Heading 2')
for r in retire: P(f"{r['parked_id']} ({pres(r['source_akt_id'])}): live {', '.join(r['matching_ids'])}",'List Bullet')
d.add_paragraph(f'2.2 Writer skips: duplicates ({len(dupskips)})',style='Heading 2')
for q in dupskips: P(f"{q['id']}: {q['notes'].strip()}",'List Bullet')
d.add_paragraph(f'2.3 Checker drops ({len(drops)})',style='Heading 2')
for x in drops: P(f"{x['id']} ({x['checker']}): {x['issues'].strip()}",'List Bullet')
P()
d.add_paragraph(f'Appendix A: items proposed for go-live ({len(fin)})',style='Heading 1')
for n,q in enumerate(fin,1):
    d.add_paragraph(f"{n}. {q['id']}: {q['presentation']}",style='Heading 2')
    P(f"Label: 2026 UKMLA Adaptation II — adapted (Oct 2026)   |   Source ID: {q['source_akt_id']}   |   Difficulty: {q['difficulty']}   |   Domain: Neurology")
    vig,qq=q['stem'].split('\n\n',1)
    p=P(); lab(p,'Stem: ',vig); P(qq)
    for l in 'ABCDE':
        p=P(); r=p.add_run(f"{l}. {q['options'][l]}"); r.bold=(l==q['correct_letter'])
    p=P(); lab(p,'Key: ',q['correct_answer'])
    p=P(); lab(p,'Why correct: ',q['why_correct'])
    p=P(); lab(p,'Why others are wrong: ',q['why_wrong'])
    m=re.search(r'Source:\s*(.*)$',q['why_correct'].strip())
    p=P(); lab(p,'Citation: ',m.group(1) if m else '; '.join(q['sources']))
    if q['id'] in flag: p=P(); lab(p,'Reviewer note: ',flag[q['id']])
    P('Verdict:  ☐ APPROVE   ☐ EDIT   ☐ REJECT      Comment: ____________________')
d.add_paragraph(f'Appendix B: holds ({len(holds)})',style='Heading 1')
P('Held because the learning point is unsound or the item is a bare recall that cannot be re-angled; neither duplicates a live question.')
for q in holds: P(f"{q['id']}: HOLD: {q['notes'].strip()}",'List Bullet')
d.save('review-packs/plab1-adaptation-ii-neurology.docx')
